"""Exercise real packaged workflow code; mock only paid-service transports."""
import base64
import hashlib
import hmac
import importlib.util
import json
import os
import sys
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

RUNTIME = Path(sys.argv.pop(1)).resolve()
sys.path.insert(0, str(RUNTIME / "skills/iflytek-contract-intelligence-review/scripts"))
sys.path.insert(0, str(RUNTIME / "bridge"))
from contract import main as contract
from contract.clients.iflytek import skill_module
from contract.clients.llm_review_client import LLMReviewClient, parse_review, signed_url
from contract.clients.ocr_client import OCRClient
from contract.clients.translate_client import TranslateClient

spec = importlib.util.spec_from_file_location("bridge", RUNTIME / "bridge/bridge.py")
bridge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bridge)
render_spec = importlib.util.spec_from_file_location("render", RUNTIME / "bridge/diagram/render.py")
render = importlib.util.module_from_spec(render_spec)
render_spec.loader.exec_module(render)

TEXT = "第一条 付款。甲方应在验收后30日内支付全部款项。第二条 争议。双方协商解决争议。"
MODEL = {"summary": "付款与争议条款摘要。", "key_clauses": ["付款期限为30日"],
         "risks": [{"title": "争议解决约定需细化", "level": "medium", "evidence": "双方协商解决争议",
                    "suggestion": "请人工核对争议解决程序。"}], "suggestions": ["核对双方主体信息"]}

class Socket:
    def __init__(self, frames=None):
        content = json.dumps(MODEL, ensure_ascii=False)
        self.frames = iter(frames if frames is not None else [
            json.dumps({"header": {"code": 0}, "payload": {"choices": {
                "status": status, "text": [{"content": part}]}}}, ensure_ascii=False)
            for status, part in [(0, content[:20]), (2, content[20:])]
        ])
        self.closed = False
        self.request = None
    def send(self, value):
        self.request = json.loads(value)
    def recv(self):
        return next(self.frames, "")
    def settimeout(self, timeout):
        assert 0 < timeout <= 45
    def close(self):
        self.closed = True

class Workflows(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.directory = Path(self.stack.enter_context(tempfile.TemporaryDirectory(prefix="ifly-contract-test-")))
        self.stack.enter_context(patch.dict(os.environ, {
            "TMP": str(self.directory), "IFLY_APP_ID": "app",
            "IFLY_API_KEY": "key", "IFLY_API_SECRET": "secret"}))
        self.config = contract.Config()
        self.socket = Socket()
        self.stack.enter_context(patch("websocket.create_connection", return_value=self.socket))
        # Any unexpected network path is an immediate failure.
        self.stack.enter_context(patch("socket.create_connection", side_effect=AssertionError("Unexpected network")))

    def request(self, parameters=None, input=None):
        return {"input": input or {"text": TEXT}, "parameters": parameters or {}}

    def test_text_bridge_produces_actual_reports_and_shared_credentials(self):
        data, artifacts = bridge.review_contract(self.request())
        self.assertEqual(len(artifacts), 2)
        self.assertEqual(data["model_review"]["summary"], MODEL["summary"])
        self.assertTrue(data["human_review_required"])
        self.assertIsNone(data["extraction_quality"]["confidence"])
        self.assertTrue(any(r["source"] == "model_inference" for r in data["risks"]))
        self.assertEqual(self.socket.request["header"]["app_id"], "app")
        self.assertEqual(self.socket.request["parameter"]["chat"]["domain"], "generalv3.5")
        self.assertTrue(self.socket.closed)
        for artifact in artifacts:
            content = (self.directory / artifact["relativePath"]).read_text(encoding="utf-8")
            self.assertIn(MODEL["summary"], content)
        self.assertEqual(json.loads((self.directory / artifacts[1]["relativePath"]).read_text(encoding="utf-8")), data)

    def test_spark_signature_uses_hostname_path_and_shared_key(self):
        parts = urlsplit(signed_url(self.config))
        query = parse_qs(parts.query)
        auth = base64.b64decode(query["authorization"][0]).decode()
        origin = f"host: {parts.netloc}\ndate: {query['date'][0]}\nGET {parts.path} HTTP/1.1"
        expected = base64.b64encode(hmac.new(b"secret", origin.encode(), hashlib.sha256).digest()).decode()
        self.assertIn('api_key="key"', auth)
        self.assertIn(expected, auth)

    def test_stream_failures_and_unverified_evidence_fail_closed(self):
        for frames in [[], ["not-json"], [json.dumps({"header": {"code": 10005, "message": "secret"}})],
                       [json.dumps({"header": {"code": 0}, "payload": {"choices": {"status": 1, "text": []}}})]]:
            sock = Socket(frames)
            with patch("websocket.create_connection", return_value=sock):
                with self.assertRaisesRegex(RuntimeError, "^Contract model review failed$"):
                    LLMReviewClient(self.config).review_contract(TEXT, "standard", [])
            self.assertTrue(sock.closed)
        changed = json.loads(json.dumps(MODEL))
        changed["risks"][0]["evidence"] = "Fabricated evidence"
        with self.assertRaises(RuntimeError):
            parse_review(json.dumps(changed), TEXT)

    def test_translation_uses_actual_body_signing_and_parser(self):
        module = skill_module("translate")
        seen = []
        def http(url, body, headers, timeout):
            payload = json.loads(body)
            self.assertEqual(payload["common"]["app_id"], "app")
            self.assertIn('api_key="key"', headers["Authorization"])
            self.assertEqual(payload["business"], {"from": "cn", "to": "en"})
            raw = base64.b64decode(payload["data"]["text"])
            self.assertLessEqual(len(raw), 1800)
            seen.append(raw.decode())
            return {"code": 0, "data": {"result": {"trans_result": {"dst": "summary"}}}}
        with patch.object(module, "_http_post", side_effect=http):
            data, _ = bridge.review_contract(self.request({"needTranslation": True}))
            self.assertEqual(data["translation_summary"], "summary")
            self.assertEqual(seen, [MODEL["summary"]])
            seen.clear()
            translated = TranslateClient(self.config).translate("中" * 1500, "zh", "en")
            self.assertEqual("".join(seen), "中" * 1500)
            self.assertIsNone(translated.confidence)
        self.socket = Socket()
        with patch("websocket.create_connection", return_value=self.socket), patch.object(module, "_http_post", return_value={"code": 1}):
            with self.assertRaises(bridge.BridgeError) as caught:
                bridge.review_contract(self.request({"needTranslation": True}))
            self.assertEqual(caught.exception.code, "UPSTREAM_ERROR")

    def ocr_response(self):
        class Response:
            def json(self):
                return {"header": {"code": 0}, "payload": {"result": {
                    "text": base64.b64encode(TEXT.encode()).decode()}}}
        return Response()

    def test_image_ocr_transport_signature_and_no_automatic_fallback(self):
        module = skill_module("ocr")
        source = self.directory / "input.png"
        from PIL import Image
        Image.new("RGB", (32, 32), "white").save(source)
        def post(url, json, headers, timeout):
            query = parse_qs(urlsplit(url).query)
            self.assertEqual(query["host"], ["cbm01.cn-huabei-1.xf-yun.com"])
            auth = base64.b64decode(query["authorization"][0]).decode()
            origin = f"host: {query['host'][0]}\ndate: {query['date'][0]}\nPOST /v1/private/se75ocrbm HTTP/1.1"
            expected = base64.b64encode(hmac.new(b"secret", origin.encode(), hashlib.sha256).digest()).decode()
            self.assertIn(expected, auth)
            self.assertEqual(json["header"]["app_id"], "app")
            self.assertEqual(json["parameter"]["ocr"]["result_format"], "markdown")
            return self.ocr_response()
        with patch.object(module.requests, "post", side_effect=post):
            data, _ = bridge.review_contract(self.request({"format": "png"}, {"files": {"document": source.name}}))
            self.assertEqual(data["document_meta"]["extraction_method"], "ocr")
        with patch.object(module.requests, "post", side_effect=RuntimeError("secret")), patch.object(contract.ImageClient, "understand_image") as image:
            with self.assertRaises(bridge.BridgeError):
                bridge.review_contract(self.request({"format": "png"}, {"files": {"document": source.name}}))
            image.assert_not_called()

    def test_explicit_image_understanding_uses_real_client_protocol(self):
        module = skill_module("image")
        source = self.directory / "input.jpg"
        from PIL import Image
        Image.new("RGB", (32, 32), "white").save(source)
        def communicate(url, body):
            request = json.loads(body)
            self.assertEqual(request["header"]["app_id"], "app")
            self.assertEqual(request["payload"]["message"]["text"][0]["content_type"], "image")
            return [json.dumps({"header": {"code": 0}, "payload": {"choices": {"status": 2, "text": [{"content": TEXT}]}}})]
        with patch.object(module, "ws_communicate", side_effect=communicate):
            data, _ = bridge.review_contract(self.request(
                {"format": "jpg", "imageMethod": "understanding"}, {"files": {"document": source.name}}))
        self.assertEqual(data["document_meta"]["extraction_method"], "image_understanding_inference")

    def test_docx_paragraphs_and_tables_are_included(self):
        from docx import Document
        source = self.directory / "input.docx"
        doc = Document()
        doc.add_paragraph(TEXT)
        doc.add_table(1, 1).cell(0, 0).text = "合同金额为100元"
        doc.save(source)
        data, _ = bridge.review_contract(self.request({"format": "docx"}, {"files": {"document": source.name}}))
        sent = json.loads(self.socket.request["payload"]["message"]["text"][1]["content"])
        self.assertIn("合同金额为100元", sent["contract"])
        self.assertEqual(data["document_meta"]["extraction_method"], "docx_extract")

    def test_pdf_rasterization_limits_and_intermediate_cleanup(self):
        import pypdfium2 as pdfium
        source = self.directory / "input.pdf"
        document = pdfium.PdfDocument.new()
        document.new_page(100, 100).close()
        document.save(source)
        document.close()
        module = skill_module("ocr")
        with patch.object(module.requests, "post", return_value=self.ocr_response()):
            data, _ = bridge.review_contract(self.request({"format": "pdf"}, {"files": {"document": source.name}}))
        self.assertEqual(data["document_meta"]["extraction_method"], "ocr")
        self.assertFalse(list(self.directory.glob("contract-pages-*")))
        document = pdfium.PdfDocument.new()
        for _ in range(9):
            document.new_page(100, 100).close()
        source.unlink()
        document.save(source)
        document.close()
        with self.assertRaises(ValueError):
            OCRClient(self.config).extract_text(str(source))

    def test_input_limits_and_original_report_summary(self):
        for params, input in [({}, {"text": "x" * 4001}), ({"lang": "invalid"}, {"text": TEXT}),
                              ({"format": "doc"}, {"text": TEXT}), ({"focus": ["invalid"]}, {"text": TEXT})]:
            with self.assertRaises(bridge.BridgeError) as caught:
                bridge.review_contract(self.request(params, input))
            self.assertEqual(caught.exception.code, "INVALID_INPUT")
        from contract.report import ReportBuilder
        result = {'risks': [{'level': 'high'}]}
        self.assertEqual(ReportBuilder.generate_summary(result), ReportBuilder().generate_summary(result))
        self.assertIn('1', ReportBuilder.generate_summary(result))

    def test_restricted_html_rejects_code_resources_and_unbounded_work(self):
        render.validate_html((RUNTIME / "bridge/diagram/workflow.html").read_text(encoding="utf-8"))
        for html in ['<script>alert(1)</script>', '<div onclick="x">x</div>', '<iframe src="file:///x"></iframe>',
                     '<img src="http://127.0.0.1/x">', '<svg><use href="#x"/></svg>',
                     '<style>@import "https://x";</style>', '<div style="background:url(file:///x)">x</div>',
                     '<style>div {background:u/**/rl(x)}</style>', '<meta http-equiv="refresh" content="0;url=http://x">',
                     '<div style="background:u\\72l(x)">x</div>']:
            with self.assertRaises(ValueError, msg=html):
                render.validate_html(html)
        with self.assertRaises(ValueError):
            render.render_html("<div>test</div>", self.directory, width=1600, height=1200, scale=2, fps=25, duration_ms=5000)
        with self.assertRaises(ValueError):
            render.render_html("<div>test</div>", self.directory, fps=True)

if __name__ == "__main__":
    unittest.main()
